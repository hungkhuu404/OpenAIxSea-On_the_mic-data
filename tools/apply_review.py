"""Turn the filled label-review workbook into golden.jsonl / heldout.jsonl.

    python apply_review.py data/cases/review/label_review.xlsx [--out data/cases]

Only rows with a reviewer name are exported, with label_status=human_confirmed.
Run freeze_eval.py afterwards. Needs openpyxl.
"""
from __future__ import annotations

import argparse
import json
import sys
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

LABELS = {"green", "yellow", "red"}


def cell(v) -> str:
    return unicodedata.normalize("NFC", str(v).strip()) if v is not None else ""


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook")
    ap.add_argument("--candidates", help="candidates.jsonl (default: ../candidates.jsonl next to the workbook)")
    ap.add_argument("--out", help="output folder (default: the candidates folder)")
    args = ap.parse_args()

    wb_path = Path(args.workbook)
    cand_path = Path(args.candidates) if args.candidates else wb_path.parent.parent / "candidates.jsonl"
    out_dir = Path(args.out) if args.out else cand_path.parent
    cands = {r["case_id"]: r for r in map(json.loads, cand_path.read_text(encoding="utf-8").splitlines()) if r}

    ws = load_workbook(wb_path, data_only=True)["Duyệt nhãn"]
    header = [cell(c.value) for c in ws[1]]
    col = {name: i for i, name in enumerate(header)}
    need = ["case_id", "Tập đề xuất", "Câu (nguyên văn)", "Đồng ý? (Y/N)", "Nhãn cuối",
            "Phải ⚪ cuối? (Y/N)", "Rule cuối", "Ghi chú", "Người duyệt"]
    missing = [n for n in need if n not in col]
    if missing:
        sys.exit(f"workbook is missing columns: {missing}")

    out = {"golden": [], "heldout": []}
    errors, skipped = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        get = lambda name: cell(row[col[name]])  # noqa: E731
        cid = get("case_id")
        if not cid:
            continue
        base = cands.get(cid)
        if base is None:
            errors.append(f"{cid}: not in candidates.jsonl")
            continue
        reviewer = get("Người duyệt")
        if not reviewer:
            skipped.append(cid)
            continue
        agree = get("Đồng ý? (Y/N)").upper()
        label = get("Nhãn cuối").lower() or base["proposed_label"]
        abstain_cell = get("Phải ⚪ cuối? (Y/N)").upper()
        abstain = (abstain_cell == "Y") if abstain_cell else bool(base.get("abstain_expected"))
        rules_cell = get("Rule cuối")
        rules = [r.strip() for r in rules_cell.split(",") if r.strip()] if rules_cell else base.get("proposed_rule_ids", [])
        if agree not in {"Y", "N"}:
            errors.append(f"{cid}: 'Đồng ý?' must be Y or N")
        if agree == "N" and not (get("Nhãn cuối") or abstain_cell or rules_cell):
            errors.append(f"{cid}: marked N but no final label / abstain / rule given")
        if label not in LABELS:
            errors.append(f"{cid}: final label '{label}' not in {sorted(LABELS)}")
        split = get("Tập đề xuất")
        if split not in out:
            errors.append(f"{cid}: split '{split}' must be golden or heldout")
            continue
        text = get("Câu (nguyên văn)")
        rec = {
            "case_id": cid, "sku_id": base["sku_id"], "text": text, "label": label,
            "abstain_expected": abstain, "rule_ids": rules,
            "legal_basis": base.get("proposed_legal_basis"),
            "source_url": base["source_url"], "source_type": base["source_type"],
            "source_date": base.get("source_date"), "modality": base["modality"],
            "label_status": "human_confirmed", "labeler": reviewer,
            "proposed_label": base["proposed_label"], "agreed_with_proposal": agree == "Y",
            "text_edited": text != base["text"], "review_note": get("Ghi chú") or None,
        }
        out[split].append(rec)

    if errors:
        print("\n".join(errors))
        return 1
    for split, recs in out.items():
        path = out_dir / f"{split}.jsonl"
        path.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in recs), encoding="utf-8")
        agree_rate = sum(r["agreed_with_proposal"] for r in recs) / len(recs) if recs else 0
        print(f"{path.name}: n={len(recs)} labels={dict(Counter(r['label'] for r in recs))} "
              f"abstain={sum(r['abstain_expected'] for r in recs)} agreement_with_proposal={agree_rate:.0%}")
    if skipped:
        print(f"skipped {len(skipped)} row(s) without a reviewer: {', '.join(skipped[:10])}{' …' if len(skipped) > 10 else ''}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
